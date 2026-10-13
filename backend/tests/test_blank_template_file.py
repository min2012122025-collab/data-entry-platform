import os
import json

from openpyxl import Workbook, load_workbook


os.environ.setdefault("ENTRY_ENV", "test")
os.environ.setdefault("ENTRY_ALLOW_SQLITE", "1")
os.environ.setdefault("ENTRY_DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("ENTRY_SECRET_KEY", "test-secret-key-for-blank-template")
os.environ.setdefault("ENTRY_COOKIE_SECURE", "0")

from app.services.import_service import ImportService  # noqa: E402


class TemplateFileDatabase:
    def __init__(self, path):
        self.path = path

    def get_template_file_record(self, template_code, version=None):
        return {
            "template_file_path": str(self.path),
            "template_file_original_name": self.path.name,
            "template_file_analysis_json": json.dumps(
                {"sheets": [{"name": "Sheet1", "header_row": 1, "example_last_row": 3}]},
                ensure_ascii=False,
            ),
        }


def test_blank_template_generation_skips_read_only_merged_cells(tmp_path):
    source = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "科目"
    sheet["B1"] = "金额"
    sheet.merge_cells("A2:A3")
    sheet["A2"] = "示例"
    sheet["B2"] = 100
    sheet["B3"] = "=B2"
    workbook.save(source)

    service = ImportService(TemplateFileDatabase(source), tmp_path / "uploads")

    path, filename = service.get_template_file("cash_flow_detail_v1", "1.0", blank=True)

    assert filename == "template-空白模板.xlsx"
    blank = load_workbook(path, data_only=False)
    sheet = blank["Sheet1"]
    assert sheet["A1"].value == "科目"
    assert sheet["A2"].value is None
    assert sheet["B2"].value is None
    assert sheet["B3"].value == "=B2"
