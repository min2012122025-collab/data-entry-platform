from openpyxl import Workbook

from app.services.import_engine import ImportPipeline


def test_import_pipeline_keeps_cash_flow_rows_with_blank_dates_without_blocking(tmp_path):
    source = tmp_path / "现金流量明细表-202608.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "HDSP_现金流量明细表_050826 (3)"
    sheet["A1"] = "现金流量明细表"
    headers = ["日期", "来源", "单据编号", "单据名称", "行号", "现金流量表项"]
    for index, header in enumerate(headers, start=2):
        sheet.cell(row=5, column=index).value = header
    sheet.merge_cells("B6:B7")
    sheet["B6"] = "2026-08-01"
    sheet.merge_cells("C6:C7")
    sheet["C6"] = "GL凭证"
    sheet["D6"] = "990020260800001"
    sheet["D7"] = "990020260800001"
    sheet["E6"] = "凭证"
    sheet["E7"] = "凭证"
    sheet["F6"] = 1
    sheet["F7"] = 2
    sheet["G6"] = "C00000"
    sheet["G7"] = "非现金流量项目"
    workbook.save(source)

    template = {
        "code": "cash_flow_detail_v1",
        "header_keywords": ["日期", "来源", "现金流量表项"],
        "date_fields": ["日期"],
        "required_fields": ["日期", "来源", "现金流量表项"],
        "validation_rules": [
            {"field": "日期", "rule_type": "required", "enabled": True},
            {"field": "来源", "rule_type": "required", "enabled": True},
            {"field": "现金流量表项", "rule_type": "required", "enabled": True},
        ],
        "template_file": {
            "analysis": {
                "confirmed": True,
                "sheets": [{"name": sheet.title, "header_row": 5}],
            },
        },
    }

    parsed = ImportPipeline().run(
        source,
        "cash_flow_detail_v1",
        template=template,
        selected_report_month="2026-08",
    )

    assert len(parsed["records"]) == 2
    assert parsed["records"][0]["row_number"] == 6
    assert parsed["records"][0]["payload"]["日期"] == "2026-08-01"
    assert parsed["records"][0]["payload"]["来源"] == "GL凭证"
    assert parsed["records"][1]["row_number"] == 7
    assert parsed["records"][1]["payload"]["日期"] is None
    assert parsed["records"][1]["payload"]["来源"] is None
    assert parsed["errors"] == []


def test_import_pipeline_matches_confirmed_single_sheet_by_headers_when_name_changes(tmp_path):
    source = tmp_path / "现金流量明细表-202608.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "HDSP_现金流量明细表_260826 (1)"
    headers = ["日期", "来源", "单据编号", "现金流量表项"]
    for index, header in enumerate(headers, start=2):
        sheet.cell(row=5, column=index).value = header
    sheet["B6"] = "2026-08-01"
    sheet["C6"] = "GL凭证"
    sheet["D6"] = "990020260800001"
    sheet["E6"] = "C00000"
    workbook.save(source)

    template = {
        "code": "cash_flow_detail_v1",
        "header_keywords": ["日期", "来源", "现金流量表项"],
        "date_fields": ["日期"],
        "required_fields": ["日期", "来源", "现金流量表项"],
        "validation_rules": [
            {"field": "日期", "rule_type": "required", "enabled": True},
            {"field": "来源", "rule_type": "required", "enabled": True},
            {"field": "现金流量表项", "rule_type": "required", "enabled": True},
        ],
        "template_file": {
            "analysis": {
                "confirmed": True,
                "sheets": [
                    {
                        "name": "HDSP_现金流量明细表_050826 (3)",
                        "header_row": 5,
                        "fields": headers,
                    }
                ],
            },
        },
    }

    parsed = ImportPipeline().run(
        source,
        "cash_flow_detail_v1",
        template=template,
        selected_report_month="2026-08",
    )

    assert parsed["records"][0]["sheet_name"] == "HDSP_现金流量明细表_050826 (3)"
    assert parsed["records"][0]["payload"]["日期"] == "2026-08-01"
    assert parsed["errors"] == []
