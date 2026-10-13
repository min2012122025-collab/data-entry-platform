from datetime import date, datetime

from openpyxl import load_workbook

from .templates import get_template


class ExcelParseError(ValueError):
    pass


def parse_monthly_report(path, template_code):
    template = get_template(template_code)
    if not template:
        raise ExcelParseError(f"Unknown template: {template_code}")

    workbook = load_workbook(path, data_only=True)
    if template["sheet"] not in workbook.sheetnames:
        raise ExcelParseError(f"Missing sheet: {template['sheet']}")

    sheet = workbook[template["sheet"]]
    report_month = _normalize_month(sheet[template["month_cell"]].value)
    if not report_month:
        raise ExcelParseError(f"Invalid report month in {template['month_cell']}")

    records = []
    errors = []
    row = template["data_start_row"]
    empty_rows = 0
    while row <= sheet.max_row:
        parsed = {}
        has_any_value = False
        for field, column in template["columns"].items():
            value = sheet[f"{column}{row}"].value
            if value not in (None, ""):
                has_any_value = True
            parsed[field] = value

        if not has_any_value:
            empty_rows += 1
            if empty_rows >= 2:
                break
            row += 1
            continue

        record = {
            "report_month": report_month,
            "row_number": row,
            "city": _clean_text(parsed["city"]),
            "metric": _clean_text(parsed["metric"]),
            "value": parsed["value"],
            "unit": _clean_text(parsed["unit"]),
            "remark": _clean_text(parsed["remark"]) or "",
        }
        records.append(record)
        errors.extend(_validate_record(record, template))
        row += 1

    if not records:
        errors.append(
            {
                "row_number": template["data_start_row"],
                "cell": f"A{template['data_start_row']}",
                "field": "city",
                "code": "empty_report",
                "message": "报表没有可导入的明细行",
                "suggestion": "请从第 5 行开始填写城市、指标、数值和单位。",
                "blocking": True,
            }
        )

    return {"report_month": report_month, "records": records, "errors": errors}


def parse_tabular_workbook(path, template_code):
    template = get_template(template_code)
    if not template:
        raise ExcelParseError(f"Unknown template: {template_code}")

    workbook = load_workbook(path, data_only=True)
    records = []
    errors = []
    report_month = None

    for sheet in workbook.worksheets:
        header_row = _find_header_row(sheet, template["header_keywords"])
        if not header_row:
            continue

        headers = _read_headers(sheet, header_row)
        for row_number in range(header_row + 1, sheet.max_row + 1):
            payload = {}
            has_value = False
            for column_index, header in headers:
                value = sheet.cell(row=row_number, column=column_index).value
                normalized = _normalize_cell_value(value)
                if normalized not in (None, ""):
                    has_value = True
                payload[header] = normalized

            if not has_value:
                continue

            if not report_month:
                report_month = _extract_report_month(payload, template.get("date_fields", []))

            records.append(
                {
                    "report_month": report_month,
                    "sheet_name": sheet.title,
                    "row_number": row_number,
                    "payload": payload,
                }
            )

    if not records:
        errors.append(
            {
                "row_number": None,
                "cell": None,
                "field": None,
                "code": "empty_report",
                "message": "没有找到可导入的表格明细",
                "suggestion": "请确认文件是当前模板，且表头未被删除或重命名。",
                "blocking": True,
            }
        )

    if not report_month:
        report_month = _extract_month_from_filename(path)

    if report_month:
        for record in records:
            record["report_month"] = report_month
    elif records:
        errors.append(
            {
                "row_number": None,
                "cell": None,
                "field": "report_month",
                "code": "missing_report_month",
                "message": "无法识别报表月份",
                "suggestion": "请确认日期字段或文件名包含年月信息。",
                "blocking": True,
            }
        )

    return {"report_month": report_month, "records": records, "errors": errors}


def _validate_record(record, template):
    errors = []
    column_by_field = template["columns"]
    labels = {"city": "城市", "metric": "指标", "value": "数值", "unit": "单位"}
    for field in template["required_fields"]:
        if record[field] in (None, ""):
            errors.append(
                {
                    "row_number": record["row_number"],
                    "cell": f"{column_by_field[field]}{record['row_number']}",
                    "field": field,
                    "code": "required",
                    "message": f"{labels[field]}不能为空",
                    "suggestion": f"请补充{labels[field]}。",
                    "blocking": True,
                }
            )

    if record["value"] not in (None, ""):
        try:
            record["value"] = float(record["value"])
        except (TypeError, ValueError):
            errors.append(
                {
                    "row_number": record["row_number"],
                    "cell": f"{column_by_field['value']}{record['row_number']}",
                    "field": "value",
                    "code": "invalid_number",
                    "message": "数值必须是数字",
                    "suggestion": "请删除文本、单位或特殊符号，只保留数字。",
                    "blocking": True,
                }
            )

    return errors


def _normalize_month(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    if value is None:
        return None
    text = str(value).strip().replace("/", "-").replace(".", "-")
    if len(text) >= 7:
        candidate = text[:7]
        try:
            datetime.strptime(candidate, "%Y-%m")
            return candidate
        except ValueError:
            return None
    return None


def _find_header_row(sheet, keywords):
    for row_number in range(1, min(sheet.max_row, 20) + 1):
        row_values = [
            _clean_text(sheet.cell(row=row_number, column=column).value)
            for column in range(1, sheet.max_column + 1)
        ]
        if all(any(keyword in (value or "") for value in row_values) for keyword in keywords):
            return row_number
    return None


def _read_headers(sheet, row_number):
    headers = []
    seen = {}
    for column in range(1, sheet.max_column + 1):
        raw = _clean_text(sheet.cell(row=row_number, column=column).value)
        if not raw:
            continue
        header = raw
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[raw]}"
        else:
            seen[header] = 1
        headers.append((column, header))
    return headers


def _normalize_cell_value(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    return value


def _extract_report_month(payload, date_fields):
    for field in date_fields:
        for key, value in payload.items():
            if field in key:
                month = _normalize_month(value)
                if month:
                    return month
    return None


def _extract_month_from_filename(path):
    import re

    name = str(path)
    match = re.search(r"(20\d{2})[.\-年]?(0?[1-9]|1[0-2])", name)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}"
    match = re.search(r"([1-9]|1[0-2])月", name)
    if match:
        return f"2026-{int(match.group(1)):02d}"
    return None


def _clean_text(value):
    if value is None:
        return None
    return str(value).strip()
