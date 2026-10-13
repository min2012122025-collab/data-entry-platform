import hashlib
import json
import shutil
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .import_engine import ImportEngineError, ImportPipeline, group_validation_errors
from .templates import SHEET_REPORT_MONTH_OFFSETS, TEMPLATES, detect_template, get_template


def offset_report_month(report_month, offset):
    year, month = (int(part) for part in report_month.split("-", 1))
    zero_based = year * 12 + month - 1 + offset
    return f"{zero_based // 12:04d}-{zero_based % 12 + 1:02d}"


class ImportService:
    def __init__(self, database, upload_dir, source_dir=None):
        self.database = database
        self.upload_dir = Path(upload_dir)
        self.source_dir = Path(source_dir) if source_dir else None
        self.pipeline = ImportPipeline()
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def list_templates(self, include_archived=False, report_month=None):
        return self.database.list_templates(include_archived=include_archived, report_month=report_month)

    def get_template(self, template_code, version=None, report_month=None):
        template = self.database.get_template_by_code(template_code, version=version, report_month=report_month)
        if not template:
            raise ValueError("模板不存在")
        return template

    def upload_template_file(self, current_user, template_code, version, source_path, filename):
        suffix = Path(filename).suffix.lower()
        if suffix not in {".xlsx", ".xlsm"}:
            raise ValueError("模板文件仅支持 .xlsx 或 .xlsm")
        size_bytes = Path(source_path).stat().st_size
        if size_bytes > 20 * 1024 * 1024:
            raise ValueError("模板文件不能超过 20MB")
        digest = self._sha256(source_path)
        target_dir = self.upload_dir / "template-files" / template_code / version
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{digest}{suffix}"
        created_target = not target.exists()
        if created_target:
            shutil.copyfile(source_path, target)
        try:
            analysis = self.analyze_template_file(target)
            return self.database.attach_template_file(
                current_user["id"],
                template_code,
                version,
                {
                    "original_name": Path(filename).name,
                    "path": target,
                    "sha256": digest,
                    "size_bytes": size_bytes,
                    "analysis": analysis,
                },
            )
        except Exception:
            if created_target:
                target.unlink(missing_ok=True)
            raise

    def analyze_template_file(self, path, header_rows=None):
        path = Path(path)
        try:
            workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
        except Exception as exc:
            raise ValueError("无法读取模板文件，请确认文件没有损坏") from exc
        sheets = []
        for sheet in workbook.worksheets:
            if header_rows is not None:
                if sheet.title not in header_rows:
                    continue
                header_row = int(header_rows[sheet.title])
                if header_row < 1 or header_row > min(sheet.max_row, 20):
                    raise ValueError(f"工作表 {sheet.title} 的表头行必须在前 20 行内")
                values = [sheet.cell(row=header_row, column=column).value for column in range(1, sheet.max_column + 1)]
                if len([value for value in values if value not in (None, "")]) < 2:
                    raise ValueError(f"工作表 {sheet.title} 第 {header_row} 行不足两个字段")
            else:
                candidates = []
                for row_number in range(1, min(sheet.max_row, 20) + 1):
                    values = [sheet.cell(row=row_number, column=column).value for column in range(1, sheet.max_column + 1)]
                    nonempty = [value for value in values if value not in (None, "")]
                    if len(nonempty) >= 2:
                        candidates.append((len(nonempty), row_number, values))
                if not candidates:
                    continue
                _, header_row, values = max(candidates, key=lambda item: (item[0], -item[1]))
            headers = []
            seen = {}
            for column, value in enumerate(values, start=1):
                if value in (None, ""):
                    continue
                raw = str(value).strip()
                seen[raw] = seen.get(raw, 0) + 1
                header = raw if seen[raw] == 1 else f"{raw}_{seen[raw]}"
                headers.append({"column": column, "name": header})
            preview_rows = []
            last_data_row = header_row
            for row_number in range(header_row + 1, sheet.max_row + 1):
                payload = {item["name"]: self._json_cell(sheet.cell(row=row_number, column=item["column"]).value) for item in headers}
                if any(value not in (None, "") for value in payload.values()):
                    last_data_row = row_number
                    if len(preview_rows) < 5:
                        preview_rows.append({"row_number": row_number, "payload": payload})
            sheets.append(
                {
                    "name": sheet.title,
                    "header_row": header_row,
                    "fields": [item["name"] for item in headers],
                    "columns": headers,
                    "example_last_row": last_data_row,
                    "preview_rows": preview_rows,
                }
            )
        if not sheets:
            raise ValueError("无法识别模板表头，请确保前 20 行中至少有一行包含两个字段")
        return {
            "sheets": sheets,
            "fields": list(dict.fromkeys(field for sheet in sheets for field in sheet["fields"])),
            "confirmed": header_rows is not None,
        }

    def confirm_template_file_headers(self, current_user, template_code, version, sheet_headers):
        record = self.database.get_template_file_record(template_code, version)
        if not record or not record.get("template_file_path"):
            raise ValueError("模板文件未上传")
        analysis = self.analyze_template_file(record["template_file_path"], dict(sheet_headers))
        return self.database.update_template_file_analysis(current_user["id"], template_code, version, analysis)

    def get_template_file(self, template_code, version=None, blank=True):
        record = self.database.get_template_file_record(template_code, version)
        if not record or not record.get("template_file_path"):
            raise ValueError("模板文件未上传")
        source = Path(record["template_file_path"])
        if not source.is_file():
            raise ValueError("模板文件已丢失，请联系管理员")
        if not blank:
            return source, record["template_file_original_name"]
        target = source.with_name(f"{source.stem}.blank{source.suffix}")
        if not target.exists():
            workbook = load_workbook(source, data_only=False, keep_vba=source.suffix.lower() == ".xlsm")
            analysis = json.loads(record.get("template_file_analysis_json") or "{}")
            for sheet_info in analysis.get("sheets", []):
                if sheet_info["name"] not in workbook.sheetnames:
                    continue
                sheet = workbook[sheet_info["name"]]
                for row in range(sheet_info["header_row"] + 1, sheet_info.get("example_last_row", sheet_info["header_row"]) + 1):
                    for column in range(1, sheet.max_column + 1):
                        cell = sheet.cell(row=row, column=column)
                        if isinstance(cell, MergedCell):
                            continue
                        if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                            cell.value = None
            workbook.save(target)
        name = f"{Path(record['template_file_original_name']).stem}-空白模板{source.suffix}"
        return target, name

    def get_template_file_preview(self, template_code, version=None):
        record = self.database.get_template_file_record(template_code, version)
        if not record or not record.get("template_file_path"):
            raise ValueError("模板文件未上传")
        return {
            "version": record["version"],
            "effective_month": record["effective_month"],
            "original_filename": record["template_file_original_name"],
            "analysis": json.loads(record.get("template_file_analysis_json") or "{}"),
        }

    def get_template_version_source_fields(self, template_code, version, source_sheet_name=None):
        record = self.database.get_template_file_record(template_code, version)
        if not record:
            raise ValueError("模板版本不存在")
        analysis = json.loads(record.get("template_file_analysis_json") or "{}")
        if source_sheet_name:
            sheet = next((item for item in analysis.get("sheets") or [] if item.get("name") == source_sheet_name), None)
            if not sheet:
                raise ValueError("模板版本中不存在该工作表")
            return sheet.get("fields") or []
        return analysis.get("fields") or []

    def _json_cell(self, value):
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return value

    def list_source_files(self):
        if not self.source_dir or not self.source_dir.exists():
            return []
        files = []
        for path in sorted(self.source_dir.glob("*.xlsx")):
            template = detect_template(path.name)
            files.append(
                {
                    "name": path.name,
                    "size": path.stat().st_size,
                    "template_code": template["code"] if template else None,
                    "template_name": template["name"] if template else "未识别模板",
                    "supported": template is not None,
                }
            )
        return files

    def sync_builtin_template_files(self, current_user, strict=False):
        if not self.source_dir:
            raise ValueError("未配置内置模板目录")
        synced = []
        missing = []
        for template in TEMPLATES:
            record = self.database.get_template_file_record(template["code"], template["version"])
            analysis = json.loads((record or {}).get("template_file_analysis_json") or "{}")
            if record and record.get("template_file_path") and analysis.get("fields") and analysis.get("confirmed"):
                continue

            source = self.source_dir / template["sample_filename"]
            if not source.is_file():
                missing.append(template["sample_filename"])
                continue

            if not record or not record.get("template_file_path"):
                uploaded = self.upload_template_file(
                    current_user,
                    template["code"],
                    template["version"],
                    source,
                    template["sample_filename"],
                )
                analysis = uploaded["template_file"]["analysis"]
            else:
                stored = Path(record["template_file_path"])
                if not stored.is_file():
                    raise ValueError(f"内置模板文件已丢失：{template['sample_filename']}")
                analysis = self.analyze_template_file(stored)

            header_rows = {sheet["name"]: sheet["header_row"] for sheet in analysis.get("sheets", [])}
            self.confirm_template_file_headers(
                current_user,
                template["code"],
                template["version"],
                header_rows,
            )
            synced.append(template["code"])

        if strict and missing:
            raise ValueError("缺少内置模板文件：" + "、".join(missing))
        return {"synced": synced, "missing": missing}

    def import_source_file(self, current_user, filename):
        if not self.source_dir:
            raise ValueError("未配置 source 目录")
        source_path = self.source_dir / Path(filename).name
        if not source_path.exists():
            raise ValueError("source 文件不存在")
        template = detect_template(source_path.name)
        if not template:
            raise ValueError("无法识别该 source 文件对应的模板")
        return self.create_import(current_user, source_path, source_path.name, template["code"])

    def create_import(self, current_user, source_path, filename, template_code, report_month=None):
        template = self.database.get_template_by_code(template_code)
        if not template:
            raise ValueError("模板不存在")
        if not report_month and current_user["role"] != "admin":
            raise PermissionError("非管理员必须从指定月份的任务上传")
        if report_month:
            self.database.require_task_write_access(current_user, template_code, report_month)

        digest = self._sha256(source_path)
        stored_path = self._store_file(source_path, filename, digest)
        raw_file = {
            "filename": filename,
            "stored_path": stored_path,
            "sha256": digest,
            "size_bytes": stored_path.stat().st_size,
            "metadata": {"content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        }

        selected_month = report_month
        try:
            parsed = self.pipeline.run(
                stored_path,
                template_code,
                template=template,
                source_name=filename,
                selected_report_month=report_month,
            )
            records = parsed["records"]
            errors = parsed["errors"]
            parsed_month = parsed["report_month"]
            selected_month = report_month or parsed_month
            if report_month:
                month_errors = []
                content_months = parsed.get("content_report_months") or []
                content_months_by_sheet = parsed.get("content_report_months_by_sheet") or {}
                sheet_offsets = SHEET_REPORT_MONTH_OFFSETS.get(template_code)
                if sheet_offsets:
                    for sheet_name, offset in sheet_offsets.items():
                        expected_month = offset_report_month(report_month, offset)
                        mismatched = [month for month in content_months_by_sheet.get(sheet_name, []) if month != expected_month]
                        if mismatched:
                            month_errors.append({
                                "row_number": None,
                                "cell": None,
                                "field": "report_month",
                                "code": "sheet_content_report_month_mismatch",
                                "message": f"{sheet_name}工作表日期期间为 {', '.join(mismatched)}，应为 {expected_month}",
                                "suggestion": f"请将{sheet_name}工作表的年月修正为 {expected_month} 后重新上传。",
                                "blocking": True,
                            })
                else:
                    mismatched_content_months = [month for month in content_months if month != report_month]
                    if mismatched_content_months:
                        month_errors.append({
                            "row_number": None,
                            "cell": None,
                            "field": "report_month",
                            "code": "content_report_month_mismatch",
                            "message": f"表内日期识别期间为 {', '.join(mismatched_content_months)}，与当前选择的 {report_month} 不一致",
                            "suggestion": "请修正表内日期，或切换到对应月份后上传。",
                            "blocking": True,
                        })
                errors = [*month_errors, *errors]
            summary = {
                **parsed["report"]["summary"],
                "_preview_fields": parsed["report"]["fields"],
                "_preview_sheets": list(dict.fromkeys(record["sheet_name"] for record in records)),
                "_preview_fields_by_sheet": {
                    sheet_name: list(dict.fromkeys(
                        field
                        for record in records if record["sheet_name"] == sheet_name
                        for field in record["payload"].keys()
                    ))
                    for sheet_name in dict.fromkeys(record["sheet_name"] for record in records)
                },
            }
        except ImportEngineError as exc:
            records = []
            errors = [
                {
                    "row_number": None,
                    "cell": None,
                    "field": None,
                    "code": "parse_error",
                    "message": str(exc),
                    "suggestion": "请确认文件使用当前模板，且表头、日期或月份信息没有被修改。",
                    "blocking": True,
                }
            ]
            summary = {
                "total_rows": 0,
                "valid_rows": 0,
                "blocking_errors": 1,
                "warning_errors": 0,
                "_preview_fields": [],
                "_preview_sheets": [],
                "_preview_fields_by_sheet": {},
            }

        return self.database.create_import_batch(
            current_user["id"],
            current_user["role"],
            template_code,
            selected_month,
            raw_file,
            summary,
            errors,
            records,
        )

    def commit_import(self, current_user, batch_id, report_month=None):
        batch = self.database.get_batch(current_user, batch_id)
        if not batch:
            raise ValueError("批次不存在")
        owner = batch.get("owner")
        if current_user["role"] != "admin" and (not owner or owner.get("id") != current_user["id"]):
            raise PermissionError("无权入库该记录")
        return self.database.commit_batch(current_user["id"], batch_id, report_month=report_month)

    def void_import(self, current_user, batch_id, reason, report_month=None):
        batch = self.database.get_batch(current_user, batch_id)
        if not batch:
            raise ValueError("批次不存在")
        owner = batch.get("owner")
        if current_user["role"] != "admin" and (not owner or owner.get("id") != current_user["id"]):
            raise PermissionError("只能作废自己负责的已入库记录")
        return self.database.void_batch(current_user["id"], batch_id, reason, report_month=report_month)

    def list_imports(self, current_user, report_month=None, status=None):
        return self.database.list_batches(current_user, report_month=report_month, status=status)

    def get_import(self, current_user, batch_id):
        batch = self.database.get_batch(current_user, batch_id)
        if not batch:
            raise ValueError("批次不存在")
        return batch

    def get_import_file(self, current_user, batch_id):
        raw_file = self.database.get_raw_file_for_batch(current_user, batch_id)
        if not raw_file:
            raise ValueError("原文件不存在")
        path = Path(raw_file["stored_path"])
        if not path.is_file():
            raise ValueError("原文件已丢失，请联系管理员检查文件存储")
        return path, raw_file["original_filename"]

    def preview_template_sample(self, template_code, page=1, page_size=20):
        template = self.get_template(template_code)
        sample = template.get("sample_filename")
        path = self.source_dir / Path(sample).name if self.source_dir and sample else None
        if not path or not path.is_file():
            raise ValueError("样例文件不存在")
        parsed = self.pipeline.run(path, template_code, template=template, source_name=path.name)
        report = parsed["report"]
        records = parsed["records"]
        total_rows = len(records)
        total_pages = (total_rows + page_size - 1) // page_size if total_rows else 0
        offset = (page - 1) * page_size
        return {
            "report_month": report["report_month"],
            "summary": report["summary"],
            "fields": report["fields"],
            "sheet_names": list(dict.fromkeys(record["sheet_name"] for record in records)),
            "rows": records[offset : offset + page_size],
            "pagination": {
                "page": page,
                "page_size": page_size,
                "total_rows": total_rows,
                "total_pages": total_pages,
            },
        }

    def get_import_preview(self, current_user, batch_id, page=1, page_size=20, sheet_name=None, limit=None):
        # ``limit`` is retained for existing internal callers while the HTTP API
        # uses explicit page/page_size parameters.
        if limit is not None:
            page = 1
            page_size = limit
        preview = self.database.get_staging_preview(current_user, batch_id, page=page, page_size=page_size, sheet_name=sheet_name)
        if not preview:
            raise ValueError("批次不存在")
        return preview

    def get_validation_report(self, current_user, batch_id):
        batch = self.get_import(current_user, batch_id)
        return group_validation_errors(batch["errors"], batch["summary"])

    def _store_file(self, source_path, filename, digest):
        source_path = Path(source_path)
        suffix = Path(filename).suffix.lower()
        target_dir = self.upload_dir / digest[:2]
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{digest}{suffix}"
        if not target.exists():
            shutil.copyfile(source_path, target)
        return target

    def _sha256(self, path):
        digest = hashlib.sha256()
        with Path(path).open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()
