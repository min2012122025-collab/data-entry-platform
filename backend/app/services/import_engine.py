from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .templates import get_template


class ImportEngineError(ValueError):
    pass


class ExcelLoader:
    def load(self, path, template):
        workbook = load_workbook(path, data_only=True)
        sheets = []
        analysis = (template.get("template_file") or {}).get("analysis") or {}
        confirmed_sheets = [
            item
            for item in analysis.get("sheets", [])
            if item.get("name") and item.get("header_row")
        ] if analysis.get("confirmed") else None
        for sheet in workbook.worksheets:
            merged_cells = self._merged_cell_lookup(sheet)
            canonical_sheet_name = sheet.title
            if confirmed_sheets is not None:
                confirmed_sheet = self._match_confirmed_sheet(
                    sheet,
                    confirmed_sheets,
                    template["header_keywords"],
                    merged_cells,
                )
                header_row = int(confirmed_sheet["header_row"]) if confirmed_sheet else None
                canonical_sheet_name = confirmed_sheet["name"] if confirmed_sheet else sheet.title
            else:
                header_row = self._find_header_row(sheet, template["header_keywords"], merged_cells)
            if not header_row:
                continue
            headers = self._read_headers(sheet, header_row, merged_cells)
            rows = []
            for row_number in range(header_row + 1, sheet.max_row + 1):
                payload = {}
                for column_index, header in headers:
                    payload[header] = sheet.cell(row=row_number, column=column_index).value
                rows.append({"sheet_name": canonical_sheet_name, "row_number": row_number, "payload": payload})
            sheets.append({"name": canonical_sheet_name, "header_row": header_row, "headers": [item[1] for item in headers], "rows": rows})
        return sheets

    def _match_confirmed_sheet(self, sheet, confirmed_sheets, keywords, merged_cells):
        exact = next((item for item in confirmed_sheets if item["name"] == sheet.title), None)
        if exact:
            return exact
        for item in confirmed_sheets:
            header_row = int(item["header_row"])
            headers = [name for _column, name in self._read_headers(sheet, header_row, merged_cells)]
            expected_fields = item.get("fields") or []
            if expected_fields and all(field in headers for field in expected_fields):
                return item
            if not expected_fields and all(any(keyword in header for header in headers) for keyword in keywords):
                return item
        return None

    def _merged_cell_lookup(self, sheet):
        lookup = {}
        for merged_range in sheet.merged_cells.ranges:
            anchor = (merged_range.min_row, merged_range.min_col)
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for column in range(merged_range.min_col, merged_range.max_col + 1):
                    if (row, column) != anchor:
                        lookup[(row, column)] = anchor
        return lookup

    def _cell_value(self, sheet, row_number, column_index, merged_cells=None):
        cell = sheet.cell(row=row_number, column=column_index)
        if not isinstance(cell, MergedCell):
            return cell.value
        anchor = (merged_cells or {}).get((row_number, column_index))
        if anchor:
            return sheet.cell(*anchor).value
        return None

    def _find_header_row(self, sheet, keywords, merged_cells=None):
        for row_number in range(1, min(sheet.max_row, 20) + 1):
            row_values = [
                self._clean_text(self._cell_value(sheet, row_number, column, merged_cells))
                for column in range(1, sheet.max_column + 1)
            ]
            if all(any(keyword in (value or "") for value in row_values) for keyword in keywords):
                return row_number
        return None

    def _read_headers(self, sheet, row_number, merged_cells=None):
        headers = []
        seen = {}
        for column in range(1, sheet.max_column + 1):
            raw = self._clean_text(self._cell_value(sheet, row_number, column, merged_cells))
            if not raw:
                continue
            header = raw
            if header in seen:
                seen[header] += 1
                header = f"{header}_{seen[raw]}"
            else:
                seen[header] = 1
            headers.append((column, header))
        return headers

    def _clean_text(self, value):
        if value is None:
            return None
        return str(value).strip()


class ExcelCleaner:
    def clean(self, loaded_sheets):
        cleaned_sheets = []
        for sheet in loaded_sheets:
            rows = []
            for row in sheet["rows"]:
                payload = {key: self._normalize_value(value) for key, value in row["payload"].items()}
                if any(value not in (None, "") for value in payload.values()):
                    rows.append({**row, "payload": payload})
            cleaned_sheets.append({**sheet, "rows": rows})
        return cleaned_sheets

    def _normalize_value(self, value):
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, str):
            value = value.strip()
            return value if value else None
        return value


class ExcelValidator:
    def validate(self, sheets, template, report_month):
        errors = []
        all_rows = [row for sheet in sheets for row in sheet["rows"]]
        for sheet in sheets:
            if not sheet["rows"]:
                errors.append(self._error(None, None, "empty_sheet", f"{sheet['name']} 没有可导入明细", "请检查该工作表是否有数据行。", True))

        rules = [rule for rule in template.get("validation_rules", []) if rule.get("enabled", True)]
        for rule_index, rule in enumerate(rules):
            errors.extend(self._validate_rule(all_rows, rule, rule_index, template.get("code")))

        if not report_month:
            errors.append(self._error(None, "report_month", "missing_report_month", "无法识别报表月份", "请确认日期字段或文件名包含年月信息。", True))
        if not sheets:
            errors.append(self._error(None, None, "missing_header", "没有找到匹配表头", "请确认文件使用当前模板。", True))
        return errors

    def _validate_rule(self, rows, rule, rule_index, template_code=None):
        field = str(rule.get("field") or "").strip()
        rule_type = rule.get("rule_type")
        parameters = rule.get("parameters") or {}
        blocking = rule.get("severity", "blocking") == "blocking"
        if not field or rule_type not in {"required", "data_type", "date", "allowed_values", "numeric_range", "unique"}:
            return []
        present = any(self._find_payload_key(row["payload"], field) for row in rows)
        if rows and not present:
            return [self._rule_error(None, field, rule, f"文件中缺少字段“{field}”", f"请使用包含“{field}”的模板文件。", blocking, rule_index)]

        errors = []
        seen = {}
        for row in rows:
            if template_code == "cash_flow_detail_v1" and row["payload"].get("日期") in (None, ""):
                continue
            key = self._find_payload_key(row["payload"], field)
            value = row["payload"].get(key) if key else None
            invalid = False
            default_message = ""
            default_suggestion = ""
            if rule_type == "required":
                invalid = value in (None, "")
                default_message = f"{field}不能为空"
                default_suggestion = f"请补充{field}后重新上传。"
            elif value in (None, ""):
                continue
            elif rule_type == "data_type":
                expected = parameters.get("type", "number")
                invalid = not self._matches_type(value, expected)
                default_message = f"{field}不是有效的{'数字' if expected == 'number' else '文本'}"
                default_suggestion = f"请按{field}要求的格式修改。"
            elif rule_type == "date":
                invalid = not self._is_date(value)
                default_message = f"{field}不是有效日期"
                default_suggestion = "请使用标准日期格式，例如 2026-07-31。"
            elif rule_type == "allowed_values":
                allowed = [str(item) for item in parameters.get("values", [])]
                invalid = str(value) not in allowed
                default_message = f"{field}不在允许值范围内"
                default_suggestion = f"请填写以下值之一：{'、'.join(allowed)}。"
            elif rule_type == "numeric_range":
                invalid = not self._is_number(value)
                if not invalid:
                    number = float(value)
                    minimum = parameters.get("min")
                    maximum = parameters.get("max")
                    invalid = (minimum not in (None, "") and number < float(minimum)) or (
                        maximum not in (None, "") and number > float(maximum)
                    )
                default_message = f"{field}超出允许范围"
                default_suggestion = "请检查该数值是否符合模板配置的范围。"
            elif rule_type == "unique":
                marker = str(value)
                invalid = marker in seen
                default_message = f"{field}存在重复值"
                default_suggestion = f"请检查第 {seen.get(marker)} 行与本行的{field}。"
                seen.setdefault(marker, row["row_number"])
            if invalid:
                errors.append(
                    self._rule_error(
                        row["row_number"],
                        key or field,
                        rule,
                        default_message,
                        default_suggestion,
                        blocking,
                        rule_index,
                    )
                )
        return errors

    def _rule_error(self, row_number, field, rule, default_message, default_suggestion, blocking, rule_index):
        return self._error(
            row_number,
            field,
            f"rule_{rule_index}_{rule.get('rule_type')}",
            rule.get("message") or default_message,
            rule.get("suggestion") or default_suggestion,
            blocking,
        )

    def _find_payload_key(self, payload, field):
        for key in payload:
            if field in key:
                return key
        return None

    def _is_number(self, value):
        try:
            float(value)
            return True
        except (TypeError, ValueError):
            return False

    def _matches_type(self, value, expected):
        if expected == "number":
            return self._is_number(value)
        if expected == "text":
            return isinstance(value, str)
        return True

    def _is_date(self, value):
        if isinstance(value, (date, datetime)):
            return True
        text = str(value).strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m", "%Y%m"):
            try:
                datetime.strptime(text, pattern)
                return True
            except ValueError:
                continue
        return False

    def _error(self, row_number, field, code, message, suggestion, blocking):
        return {
            "row_number": row_number,
            "cell": None,
            "field": field,
            "code": code,
            "message": message,
            "suggestion": suggestion,
            "blocking": blocking,
        }


class SchemaMapper:
    def map(self, sheets, report_month):
        records = []
        for sheet in sheets:
            for row in sheet["rows"]:
                records.append(
                    {
                        "report_month": report_month,
                        "sheet_name": sheet["name"],
                        "row_number": row["row_number"],
                        "payload": row["payload"],
                    }
                )
        return records

    def describe_fields(self, sheets):
        fields = []
        seen = set()
        for sheet in sheets:
            for header in sheet["headers"]:
                if header not in seen:
                    fields.append(header)
                    seen.add(header)
        return fields


class PreviewReport:
    def build(self, template, sheets, records, errors, fields, report_month, limit=5):
        blocking_errors = sum(1 for error in errors if error.get("blocking"))
        return {
            "report_month": report_month,
            "summary": {
                "total_rows": len(records),
                "valid_rows": len(records) if not blocking_errors else 0,
                "blocking_errors": blocking_errors,
                "warning_errors": len(errors) - blocking_errors,
            },
            "sheets": [
                {"name": sheet["name"], "header_row": sheet["header_row"], "rows": len(sheet["rows"]), "columns": sheet["headers"]}
                for sheet in sheets
            ],
            "fields": fields,
            "preview_rows": records[:limit],
            "errors": errors,
            "target_strategy": template.get("target_strategy", "unconfigured"),
        }


class ImportPipeline:
    def __init__(self):
        self.loader = ExcelLoader()
        self.cleaner = ExcelCleaner()
        self.validator = ExcelValidator()
        self.mapper = SchemaMapper()
        self.reporter = PreviewReport()

    def run(self, path, template_code, template=None, source_name=None, selected_report_month=None):
        template = template or get_template(template_code)
        if not template:
            raise ImportEngineError(f"Unknown template: {template_code}")
        sheets = self.cleaner.clean(self.loader.load(path, template))
        full_filename_month, filename_month_hint = self._extract_month_from_filename(source_name or path)
        content_report_months = self._extract_report_months(sheets, template.get("date_fields", []))
        content_report_months_by_sheet = self._extract_report_months_by_sheet(sheets, template.get("date_fields", []))
        filename_report_month = full_filename_month
        if not filename_report_month and filename_month_hint and selected_report_month:
            selected_year, _selected_month = selected_report_month.split("-", 1)
            filename_report_month = f"{selected_year}-{filename_month_hint:02d}"

        # The page/task month owns the import scope. Filename and workbook dates are
        # independent evidence that the service layer validates against that scope;
        # neither source may hide a disagreement in the other.
        report_month = selected_report_month or (content_report_months[0] if content_report_months else None) or filename_report_month
        records = self.mapper.map(sheets, report_month)
        fields = self.mapper.describe_fields(sheets)
        errors = self.validator.validate(sheets, template, report_month)
        report = self.reporter.build(template, sheets, records, errors, fields, report_month)
        return {
            "report_month": report_month,
            "filename_report_month": filename_report_month,
            "content_report_months": content_report_months,
            "content_report_months_by_sheet": content_report_months_by_sheet,
            "records": records,
            "errors": errors,
            "report": report,
        }

    def _extract_report_month(self, sheets, date_fields):
        months = self._extract_report_months(sheets, date_fields)
        return months[0] if months else None

    def _extract_report_months(self, sheets, date_fields):
        months = []
        for sheet in sheets:
            for row in sheet["rows"]:
                for field in date_fields:
                    for key, value in row["payload"].items():
                        if field in key:
                            month = self._normalize_month(value)
                            if month and month not in months:
                                months.append(month)
        return months

    def _extract_report_months_by_sheet(self, sheets, date_fields):
        return {
            sheet["name"]: self._extract_report_months([sheet], date_fields)
            for sheet in sheets
        }

    def _normalize_month(self, value):
        if isinstance(value, (datetime, date)):
            return value.strftime("%Y-%m")
        if value is None:
            return None
        text_value = str(value).strip().replace("/", "-").replace(".", "-")
        if len(text_value) >= 7:
            candidate = text_value[:7]
            try:
                datetime.strptime(candidate, "%Y-%m")
                return candidate
            except ValueError:
                return None
        return None

    def _extract_month_from_filename(self, path):
        name = str(path)
        match = re.search(r"(20\d{2})[.\-年]?(0?[1-9]|1[0-2])", name)
        if match:
            return f"{match.group(1)}-{int(match.group(2)):02d}", None
        match = re.search(r"([1-9]|1[0-2])月", name)
        if match:
            return None, int(match.group(1))
        return None, None


def group_validation_errors(errors, summary):
    grouped = defaultdict(lambda: {"count": 0, "samples": []})
    for error in errors:
        severity = "blocking" if error.get("blocking") else "warning"
        key = (severity, error["code"], error["message"])
        grouped[key]["count"] += 1
        if len(grouped[key]["samples"]) < 5:
            grouped[key]["samples"].append(error)
    return {
        "summary": summary,
        "groups": [
            {"severity": severity, "code": code, "message": message, **payload}
            for (severity, code, message), payload in grouped.items()
        ],
    }
